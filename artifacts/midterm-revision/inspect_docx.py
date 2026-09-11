import json
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8")


def emu_to_cm(value):
    return None if value is None else round(value.cm, 3)


def run_info(run):
    rpr = run._element.rPr
    rfonts = None if rpr is None else rpr.rFonts
    east_asia = None if rfonts is None else rfonts.get(qn("w:eastAsia"))
    ascii_font = None if rfonts is None else rfonts.get(qn("w:ascii"))
    return {
        "text": run.text,
        "font_name": run.font.name,
        "east_asia": east_asia,
        "ascii": ascii_font,
        "size_pt": None if run.font.size is None else round(run.font.size.pt, 2),
        "bold": run.bold,
        "italic": run.italic,
    }


def paragraph_info(index, paragraph):
    pf = paragraph.paragraph_format
    return {
        "index": index,
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": None if paragraph.alignment is None else int(paragraph.alignment),
        "left_indent_cm": emu_to_cm(pf.left_indent),
        "first_line_indent_cm": emu_to_cm(pf.first_line_indent),
        "space_before_pt": None if pf.space_before is None else round(pf.space_before.pt, 2),
        "space_after_pt": None if pf.space_after is None else round(pf.space_after.pt, 2),
        "line_spacing": pf.line_spacing,
        "runs": [run_info(r) for r in paragraph.runs if r.text],
    }


def main(path_text):
    path = Path(path_text)
    doc = Document(path)
    sections = []
    for i, s in enumerate(doc.sections, 1):
        sections.append({
            "index": i,
            "width_cm": emu_to_cm(s.page_width),
            "height_cm": emu_to_cm(s.page_height),
            "top_cm": emu_to_cm(s.top_margin),
            "bottom_cm": emu_to_cm(s.bottom_margin),
            "left_cm": emu_to_cm(s.left_margin),
            "right_cm": emu_to_cm(s.right_margin),
            "header_cm": emu_to_cm(s.header_distance),
            "footer_cm": emu_to_cm(s.footer_distance),
            "different_first_page": s.different_first_page_header_footer,
        })
    all_paras = [paragraph_info(i, p) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
    paras = []
    for p in all_paras:
        text = p["text"].strip()
        heading_like = (
            len(text) <= 90
            and (text[:1].isdigit() or "报告" in text or "研究" in text or "专　　利" in text or "标　　准" in text)
        )
        if len(paras) < 24 or heading_like:
            paras.append(p)
    signatures = Counter()
    style_usage = Counter()
    for p in all_paras:
        style_usage[p["style"]] += 1
        for run in p["runs"]:
            key = (
                run["east_asia"] or run["font_name"] or run["ascii"],
                run["size_pt"],
                run["bold"],
            )
            signatures[key] += len(run["text"])
    styles = {}
    requested_names = ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "正文", "标题 1", "标题 2", "标题 3"]
    requested_names.extend(
        st.name for st in doc.styles
        if st.type == WD_STYLE_TYPE.PARAGRAPH and ("标题" in st.name or "正文" in st.name)
    )
    for name in dict.fromkeys(requested_names):
        if name not in doc.styles:
            continue
        st = doc.styles[name]
        pf = st.paragraph_format
        styles[name] = {
            "font_name": st.font.name,
            "size_pt": None if st.font.size is None else round(st.font.size.pt, 2),
            "bold": st.font.bold,
            "left_indent_cm": emu_to_cm(pf.left_indent),
            "first_line_indent_cm": emu_to_cm(pf.first_line_indent),
            "space_before_pt": None if pf.space_before is None else round(pf.space_before.pt, 2),
            "space_after_pt": None if pf.space_after is None else round(pf.space_after.pt, 2),
            "line_spacing": pf.line_spacing,
        }
    print(json.dumps({
        "path": str(path.resolve()),
        "sections": sections,
        "styles": styles,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "top_run_signatures": [
            {"font": k[0], "size_pt": k[1], "bold": k[2], "characters": v}
            for k, v in signatures.most_common(15)
        ],
        "style_usage": dict(style_usage.most_common()),
        "paragraphs": paras,
    }, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main(sys.argv[1])
