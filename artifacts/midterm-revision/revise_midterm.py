import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt


SOURCE_DOCX = Path(r"D:\硕士毕设\docs\reports\midterm\中期报告-匡智颉.docx")
FINAL_DOCX = Path(r"D:\硕士毕设\docs\reports\midterm\中期报告-匡智颉-修订版-字体与引用更新.docx")
SOURCE_MD = Path(r"D:\硕士毕设\docs\reports\midterm\中期报告-匡智颉.md")
FINAL_MD = Path(r"D:\硕士毕设\docs\reports\midterm\中期报告-匡智颉-修订版-字体与引用更新.md")


SECTION_FOUR = """## 4 主要研究进展

开题以来，课题工作重点是收敛混合求解方案、明确正确性边界并完成基础接口实现。当前已形成可执行的系统结构和实验方案，但尚未完成可用于评价算法效果的正式对比实验，因此本章只陈述已落实的设计与实现工作，不报告实验效果。

### 4.1 研究方案收敛与正确性边界

课题已明确以加权部分 MaxSAT 为研究对象，由 CASHWMaxSAT-DisjCad-S6 承担精确推理、下界维护和最优性证明，由 SPBMAXSAT2 负责在限定时间内构造可行解并给出候选上界。局部搜索结果只有在通过原始 WCNF 语义复核并严格改善当前上界时，才允许进入精确求解器；最优性仍必须由 CASH 通过合法终止状态确认。

双向交互采用受约束的状态投影。CASH 只导出原始变量中当前仍有效的传播赋值，未确定变量记为 -1，不导出分支决策、松弛变量或辅助变量。SPB 可把该部分赋值作为搜索起点并在局部搜索中自由翻转，但不会把变量赋值写回 CASH。该边界避免了启发式状态干扰精确求解器的证明过程。

调度方案先采用固定时间片：CASH 连续运行一个时间窗口，在安全点暂停；SPB 完整运行一个较短窗口后仅返回候选上界；随后 CASH 从原搜索状态继续。解析、候选验证、模块切换和日志记录均计入统一的端到端预算，为后续比较提供一致的资源口径。

### 4.2 SPB 解析与持久化接口

针对 SPB 对标准 WCNF 的适配问题，当前实现改为先读取并校验文件头，再依据变量数和子句数分配存储；数值 top 与 h 关键字采用统一的硬子句语义。实现同时检查文字范围、子句数量和结束符，并将空软子句计入固定目标代价、将空硬子句判为不可行。相应的回归用例已经准备，需在服务器环境完成构建和执行后确认。

为支持多轮协同，SPB 新增独立的持久化调用接口。该接口在同一实例内复用工作器并保留自适应子句权重；每轮接收新的部分赋值后，重新计算满足计数、评分和候选变量栈，不保留上一轮赋值，也不改变原有局部搜索策略和权重更新公式。

SPB 封装内部保留候选模型作为数值上界的证书。候选提交前需重新检查全部硬子句，并按原始软子句权重计算代价；模型长度、硬可行性或目标值任一不一致时，候选不得提交给 CASH。该模型只用于内部校验，不参与 CASH 的变量赋值和证明过程。

### 4.3 CASH 安全点与外部上界接收

CASH 侧采用最小侵入的协同回调。回调位于目标变换完成后的优化主循环安全点，在一次求解调用及其状态处理完整结束后执行，负责导出原始变量传播快照、接收候选上界并返回调度控制。精确求解器的核心提取、下界推进和子句硬化逻辑保持不变。

外部上界通过唯一的单位转换路径进入 CASH。接口会拒绝负值、无法按内部目标单位转换的值以及低于当前下界的值；合法候选只更新正常的上界状态，不直接改写下界，也不能单独触发最优性结论。针对空软子句固定代价和权重最大公约数的处理也已纳入统一转换规则。

### 4.4 混合协调器与实验准备

项目已建立 HybridMaxSAT 协调器和统一运行器。协调器持有单个 SPB 对象，负责固定时间片调度、原变量快照截取、候选上界交接和失败回退；运行器读取 WCNF 实例并输出结构化记录，包括每轮时间、传播变量数、交接前后的上界以及候选接收状态。

实验方案设置 CASH、SPB、Hybrid 和 Hybrid-NoInference 四种配置。前两者分别作为精确算法和局部搜索基线；Hybrid 使用 CASH 的传播赋值初始化 SPB；Hybrid-NoInference 保持相同时间片但采用随机初始化，用于分析传播信息的作用。所有配置使用相同的总预算、实例清单和随机种子，并记录构建版本、配置和原始日志。

当前已完成实验入口、实例清单和日志字段的准备工作。服务器构建、回归测试和正式对比实验尚未完成，现阶段不对算法效果作结论。

### 4.5 实验进展

#### 4.5.1 当前实验效果


#### 4.5.2 已尝试的改进方法


### 4.6 已取得的阶段性成果

1. 明确了 CASH 与 SPB 的职责、状态边界、上界接收规则和最优性判据；
2. 完成 SPB 标准 WCNF 解析修复及跨轮持久化接口的代码实现，并准备了对应回归用例；
3. 完成 CASH 安全点回调、原变量传播快照导出和外部上界转换接口的代码实现；
4. 完成混合协调器、统一运行器、顶层构建入口和结构化日志设计；
5. 建立后续实验的配置矩阵、实例清单、随机种子和验收规则，待服务器验证后开展正式实验。

**发表与接收论文**

暂无论文发表或录用。

**专　　利**

暂无已申请或授权专利。

**标　　准**

无。

**科技奖励**

暂无。
"""


DOC_PARAGRAPHS = [
    ("body", "开题以来，课题工作重点是收敛混合求解方案、明确正确性边界并完成基础接口实现。当前已形成可执行的系统结构和实验方案，但尚未完成可用于评价算法效果的正式对比实验，因此本章只陈述已落实的设计与实现工作，不报告实验效果。"),
    ("h2", "4.1 研究方案收敛与正确性边界"),
    ("body", "课题已明确以加权部分 MaxSAT 为研究对象，由 CASHWMaxSAT-DisjCad-S6 承担精确推理、下界维护和最优性证明，由 SPBMAXSAT2 负责在限定时间内构造可行解并给出候选上界。局部搜索结果只有在通过原始 WCNF 语义复核并严格改善当前上界时，才允许进入精确求解器；最优性仍必须由 CASH 通过合法终止状态确认。"),
    ("body", "双向交互采用受约束的状态投影。CASH 只导出原始变量中当前仍有效的传播赋值，未确定变量记为 -1，不导出分支决策、松弛变量或辅助变量。SPB 可把该部分赋值作为搜索起点并在局部搜索中自由翻转，但不会把变量赋值写回 CASH。该边界避免了启发式状态干扰精确求解器的证明过程。"),
    ("body", "调度方案先采用固定时间片：CASH 连续运行一个时间窗口，在安全点暂停；SPB 完整运行一个较短窗口后仅返回候选上界；随后 CASH 从原搜索状态继续。解析、候选验证、模块切换和日志记录均计入统一的端到端预算，为后续比较提供一致的资源口径。"),
    ("h2", "4.2 SPB 解析与持久化接口"),
    ("body", "针对 SPB 对标准 WCNF 的适配问题，当前实现改为先读取并校验文件头，再依据变量数和子句数分配存储；数值 top 与 h 关键字采用统一的硬子句语义。实现同时检查文字范围、子句数量和结束符，并将空软子句计入固定目标代价、将空硬子句判为不可行。相应的回归用例已经准备，需在服务器环境完成构建和执行后确认。"),
    ("body", "为支持多轮协同，SPB 新增独立的持久化调用接口。该接口在同一实例内复用工作器并保留自适应子句权重；每轮接收新的部分赋值后，重新计算满足计数、评分和候选变量栈，不保留上一轮赋值，也不改变原有局部搜索策略和权重更新公式。"),
    ("body", "SPB 封装内部保留候选模型作为数值上界的证书。候选提交前需重新检查全部硬子句，并按原始软子句权重计算代价；模型长度、硬可行性或目标值任一不一致时，候选不得提交给 CASH。该模型只用于内部校验，不参与 CASH 的变量赋值和证明过程。"),
    ("h2", "4.3 CASH 安全点与外部上界接收"),
    ("body", "CASH 侧采用最小侵入的协同回调。回调位于目标变换完成后的优化主循环安全点，在一次求解调用及其状态处理完整结束后执行，负责导出原始变量传播快照、接收候选上界并返回调度控制。精确求解器的核心提取、下界推进和子句硬化逻辑保持不变。"),
    ("body", "外部上界通过唯一的单位转换路径进入 CASH。接口会拒绝负值、无法按内部目标单位转换的值以及低于当前下界的值；合法候选只更新正常的上界状态，不直接改写下界，也不能单独触发最优性结论。针对空软子句固定代价和权重最大公约数的处理也已纳入统一转换规则。"),
    ("h2", "4.4 混合协调器与实验准备"),
    ("body", "项目已建立 HybridMaxSAT 协调器和统一运行器。协调器持有单个 SPB 对象，负责固定时间片调度、原变量快照截取、候选上界交接和失败回退；运行器读取 WCNF 实例并输出结构化记录，包括每轮时间、传播变量数、交接前后的上界以及候选接收状态。"),
    ("body", "实验方案设置 CASH、SPB、Hybrid 和 Hybrid-NoInference 四种配置。前两者分别作为精确算法和局部搜索基线；Hybrid 使用 CASH 的传播赋值初始化 SPB；Hybrid-NoInference 保持相同时间片但采用随机初始化，用于分析传播信息的作用。所有配置使用相同的总预算、实例清单和随机种子，并记录构建版本、配置和原始日志。"),
    ("body", "当前已完成实验入口、实例清单和日志字段的准备工作。服务器构建、回归测试和正式对比实验尚未完成，现阶段不对算法效果作结论。"),
    ("h2", "4.5 实验进展"),
    ("h3", "4.5.1 当前实验效果"),
    ("blank", ""),
    ("h3", "4.5.2 已尝试的改进方法"),
    ("blank", ""),
    ("h2", "4.6 已取得的阶段性成果"),
    ("list", "1. 明确了 CASH 与 SPB 的职责、状态边界、上界接收规则和最优性判据；"),
    ("list", "2. 完成 SPB 标准 WCNF 解析修复及跨轮持久化接口的代码实现，并准备了对应回归用例；"),
    ("list", "3. 完成 CASH 安全点回调、原变量传播快照导出和外部上界转换接口的代码实现；"),
    ("list", "4. 完成混合协调器、统一运行器、顶层构建入口和结构化日志设计；"),
    ("list", "5. 建立后续实验的配置矩阵、实例清单、随机种子和验收规则，待服务器验证后开展正式实验。"),
    ("bold_body", "发表与接收论文"),
    ("body", "暂无论文发表或录用。"),
    ("bold_body", "专　　利"),
    ("body", "暂无已申请或授权专利。"),
    ("bold_body", "标　　准"),
    ("body", "无。"),
    ("bold_body", "科技奖励"),
    ("body", "暂无。"),
]


# 文献序号与第 6 章参考文献一一对应。仅在可由文献直接支持的研究背景、
# 研究现状和方案论述处添加，不将尚未完成的实验工作包装成引用结论。
CITATIONS = {
    "最大可满足性问题（Maximum Satisfiability": "[1, 34-36]",
    "MaxSAT 的建模能力覆盖": "[32, 33, 35, 36]",
    "求解 MaxSAT 的算法分为": "[1-20]",
    "本课题的理论意义在于": "[1, 29, 31]",
    "基于 SAT 求解器的 MaxSAT 精确": "[2, 12, 13]",
    "核心引导方法的演进可分为": "[2-10]",
    "2011 年至今为框架化": "[10-20]",
    "不完备算法以局部搜索为主": "[21-25]",
    "里程碑是 Dist": "[24-30]",
    "由于两类方法各自触及瓶颈": "[29, 31]",
    "精确算法的冷启动瓶颈未被根本": "[2-20]",
    "局部搜索解质量不稳定且无最优性": "[21-30]",
    "已有混合方法交互粒度粗": "[29, 31]",
    "两类求解器的状态空间与目标单位": "[12, 13, 33]",
}


def set_run_font(run, east_asia, ascii_font, size, bold=None):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)


def set_run_family(run, east_asia="宋体", ascii_font="Times New Roman"):
    """Set families only, retaining the original table/header size and emphasis."""
    run.font.name = ascii_font
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)


def normalize_container_families(container):
    """Cover tables and page furniture omitted by Document.paragraphs."""
    for paragraph in container.paragraphs:
        for run in paragraph.runs:
            set_run_family(run)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                normalize_container_families(cell)


def set_paragraph_format(paragraph, role):
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.5
    pf.keep_with_next = role in {"h1", "h2", "h3"}
    pf.page_break_before = False
    if role == "h1":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        font = ("宋体", "Times New Roman", 14, True)
    elif role == "h2":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(9)
        pf.space_after = Pt(4)
        font = ("宋体", "Times New Roman", 12, True)
    elif role == "h3":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(7.8)
        pf.space_after = Pt(4)
        font = ("宋体", "Times New Roman", 12, True)
    elif role == "blank":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.847)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        font = ("宋体", "Times New Roman", 12, False)
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.847)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        font = ("宋体", "Times New Roman", 12, role == "bold_body")
    for run in paragraph.runs:
        set_run_font(run, *font)


def insert_paragraph_before(reference, text, role):
    new_p = OxmlElement("w:p")
    reference._p.addprevious(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    if text:
        paragraph.add_run(text)
    set_paragraph_format(paragraph, role)
    return paragraph


def append_citation(paragraph):
    """Append the matching numeric citation once, preserving the original prose."""
    text = paragraph.text.strip()
    if not text or re.search(r"\[\d+(?:[-,]\s*\d+)*\]$", text):
        return
    for prefix, citation in CITATIONS.items():
        if text.startswith(prefix):
            paragraph.add_run(" " + citation)
            return


def replace_section_four(doc):
    body = doc._element.body
    start = None
    end = None
    for child in list(body):
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc._body)
        text = paragraph.text.strip()
        if text == "4 主要研究进展":
            start = child
        elif text == "5 后续研究计划和预期成果":
            end = child
            break
    if start is None or end is None:
        raise RuntimeError("Unable to locate section 4 boundaries")
    children = list(body)
    start_index = children.index(start)
    end_index = children.index(end)
    for child in children[start_index + 1:end_index]:
        body.remove(child)
    start_paragraph = Paragraph(start, doc._body)
    set_paragraph_format(start_paragraph, "h1")
    end_paragraph = Paragraph(end, doc._body)
    for role, text in DOC_PARAGRAPHS:
        insert_paragraph_before(end_paragraph, text, role)


def normalize_document(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.7)
        section.right_margin = Cm(2.7)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.3)
        section.different_first_page_header_footer = True

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run("融合局部搜索与核心引导的混合精确MaxSAT求解算法研究")
        set_run_font(run, "宋体", "Times New Roman", 12, False)
        ppr = hp._p.get_or_add_pPr()
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            ppr.append(pbdr)
        bottom = pbdr.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            pbdr.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")

        first_header = section.first_page_header
        first_header.is_linked_to_previous = False
        for p in first_header.paragraphs:
            p.clear()

        normalize_container_families(section.header)
        normalize_container_families(section.footer)
        normalize_container_families(section.first_page_header)
        normalize_container_families(section.first_page_footer)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text == "硕士生论文中期进展报告":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, "宋体", "Times New Roman", 22, True)
            continue
        if text.startswith("题 目："):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, "宋体", "Times New Roman", 14, True)
            continue
        if re.match(r"^[1-6]\s+\S", text):
            set_paragraph_format(paragraph, "h1")
        elif re.match(r"^\d+\.\d+\.\d+\s+\S", text):
            set_paragraph_format(paragraph, "h3")
        elif re.match(r"^\d+\.\d+\s+\S", text):
            set_paragraph_format(paragraph, "h2")
        elif paragraph._p.xpath(".//w:instrText"):
            continue
        else:
            role = "bold_body" if text in {"发表与接收论文", "专　　利", "标　　准", "科技奖励"} else "body"
            set_paragraph_format(paragraph, role)
            append_citation(paragraph)
            set_paragraph_format(paragraph, role)

    normalize_container_families(doc)

    # Ensure the explicit cover page break remains, now that the tighter margins keep the cover on one page.
    cover_breaks = [p for p in doc.paragraphs[:20] if p._p.xpath(".//w:br[@w:type='page']") or p._p.xpath(".//w:br")]
    if not cover_breaks:
        for p in doc.paragraphs:
            if p.text.strip() == "华中科技大学研究生院制":
                new_p = OxmlElement("w:p")
                p._p.addnext(new_p)
                new_paragraph = Paragraph(new_p, p._parent)
                run = new_paragraph.add_run()
                br = OxmlElement("w:br")
                br.set(qn("w:type"), "page")
                run._r.append(br)
                break


def revise_markdown():
    source = SOURCE_MD.read_text(encoding="utf-8")
    start = source.index("## 4 主要研究进展")
    end = source.index("## 5 后续研究计划和预期成果")
    revised = source[:start] + SECTION_FOUR + "\n" + source[end:]
    for prefix, citation in CITATIONS.items():
        marker = prefix
        marker_index = revised.find(marker)
        if marker_index < 0:
            continue
        paragraph_end = revised.find("\n\n", marker_index)
        if paragraph_end < 0:
            continue
        paragraph = revised[marker_index:paragraph_end]
        if not re.search(r"\[\d+(?:[-,]\s*\d+)*\]$", paragraph):
            revised = revised[:paragraph_end] + " " + citation + revised[paragraph_end:]
    FINAL_MD.write_text(revised, encoding="utf-8", newline="\n")


def main():
    shutil.copy2(SOURCE_DOCX, FINAL_DOCX)
    doc = Document(FINAL_DOCX)
    replace_section_four(doc)
    normalize_document(doc)
    doc.save(FINAL_DOCX)
    revise_markdown()
    print(FINAL_DOCX)
    print(FINAL_MD)


if __name__ == "__main__":
    main()
